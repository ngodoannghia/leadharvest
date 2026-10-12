"""
Xuất dữ liệu bàn giao sang định dạng Excel (.xlsx) chuyên nghiệp (BA §6, §9).

Yêu cầu kỹ thuật:
  1. Gồm 6 tab tương ứng 6 phân khúc mục tiêu theo thứ tự quy định.
  2. Tab thứ 7: 'Thong_Ke' hiển thị toàn bộ báo cáo chất lượng QA Report (BA §9).
  3. Định dạng text (@) tuyệt đối cho các cột SĐT và MST để bảo toàn số 0 ở đầu khi mở bằng Excel.
  4. Header phong cách chuẩn doanh nghiệp (Nền xanh đậm, chữ trắng, căn giữa, tự động căn chỉnh độ rộng cột).
"""

from __future__ import annotations

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, numbers
from openpyxl.utils import get_column_letter

from leadharvest.models.enums import SEGMENT_ORDER, Segment
from leadharvest.models.lead_output import LeadOutput
from leadharvest.processing.qa_report import QAReport


def export_to_excel(
    leads: list[LeadOutput],
    qa_report: QAReport | None,
    output_path: str,
) -> None:
    """
    Xuất danh sách LeadOutput và QAReport vào file Excel đa tab hoàn chỉnh.
    """
    wb = openpyxl.Workbook()
    # Xóa sheet mặc định ban đầu
    wb.remove(wb.active)

    # Style chuẩn cho Header bảng dữ liệu
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Phân loại leads theo từng phân khúc
    leads_by_seg: dict[str, list[LeadOutput]] = {s.value: [] for s in SEGMENT_ORDER}
    for lead in leads:
        if lead.target_segment in leads_by_seg:
            leads_by_seg[lead.target_segment].append(lead)

    # 1. Tạo 6 sheet cho 6 phân khúc ngành
    for seg in SEGMENT_ORDER:
        ws = wb.create_sheet(title=seg.value)
        seg_leads = leads_by_seg[seg.value]

        # Ghi header 11 cột
        ws.append(list(LeadOutput.COLUMN_HEADERS))
        ws.row_dimensions[1].height = 28

        for col_idx in range(1, len(LeadOutput.COLUMN_HEADERS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # Ghi từng dòng dữ liệu
        for row_idx, lead in enumerate(seg_leads, start=2):
            ws.append(list(lead.to_row()))
            ws.row_dimensions[row_idx].height = 20

            for col_idx in range(1, len(LeadOutput.COLUMN_HEADERS) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

                # Cột SĐT chính (cột 4), SĐT phụ (cột 6), MST (cột 8) phải định dạng TEXT (@)
                if col_idx in (4, 6, 8):
                    cell.number_format = numbers.FORMAT_TEXT
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Tự động điều chỉnh độ rộng cột
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                val_len = max(len(line) for line in val_str.split("\n")) if "\n" in val_str else len(val_str)
                if val_len > max_len:
                    max_len = val_len
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 60)

    # 2. Tạo tab thứ 7: Thong_Ke (QA Report) nếu có
    if qa_report:
        ws_stat = wb.create_sheet(title="Thong_Ke")
        _write_qa_summary_sheet(ws_stat, qa_report)

    wb.save(output_path)


def _write_qa_summary_sheet(ws, report: QAReport) -> None:
    """Vẽ bảng tổng hợp chất lượng dữ liệu QA Report trên sheet Thong_Ke."""
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    section_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    section_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    bold_font = Font(name="Calibri", size=11, bold=True)
    pass_font = Font(name="Calibri", size=11, bold=True, color="008000")
    fail_font = Font(name="Calibri", size=11, bold=True, color="FF0000")

    ws.append(["BÁO CÁO CHẤT LƯỢNG DỮ LIỆU BÀN GIAO (QA REPORT)"])
    ws.cell(row=1, column=1).font = title_font
    ws.append([])

    def add_section(title: str):
        row = ws.max_row + 1
        ws.cell(row=row, column=1, value=title).font = section_font
        ws.cell(row=row, column=1).fill = section_fill
        ws.cell(row=row, column=2).fill = section_fill

    # Phần 1: Sản lượng
    add_section("1. TỔNG QUAN SẢN LƯỢNG")
    ws.append(["Tổng số bản ghi thô thu thập", report.total_raw_records])
    ws.append(["Bị loại do trùng lặp (Duplicate)", report.rejected_duplicates])
    ws.append(["Bị loại do thiếu thông tin (Missing info)", report.rejected_missing_info])
    ws.append(["Bị loại do sai ngành (Wrong segment)", report.rejected_wrong_segment])
    ws.append(["Bị loại do ngừng hoạt động (Inactive)", report.rejected_inactive])
    ws.append(["Tổng số bản ghi bị loại", report.total_rejected])
    ws.append(["Tổng số bản ghi BÀN GIAO CHÍNH THỨC", report.total_delivered_records])
    ws.cell(row=ws.max_row, column=1).font = bold_font
    ws.cell(row=ws.max_row, column=2).font = bold_font
    ws.append([])

    # Phần 2: Sản lượng theo ngành
    add_section("2. SẢN LƯỢNG BÀN GIAO THEO PHÂN KHÚC")
    for seg in SEGMENT_ORDER:
        count = report.segment_counts.get(seg.value, 0)
        target = seg.min_quota
        status = "ĐẠT" if count >= target else "CHƯA ĐẠT"
        ws.append([f"{seg.value} (Sàn: {target:,})", count, status])
    ws.append([])

    # Phần 3: Thống kê SĐT
    add_section("3. THỐNG KÊ SỐ ĐIỆN THOẠI")
    ws.append(["Số dòng có số di động (Mobile)", report.total_with_mobile])
    ws.append(["Số dòng có số bàn (Landline)", report.total_with_landline])
    ws.append(["Số dòng có hotline", report.total_with_hotline])
    ws.append(["Tỷ lệ số di động (%) (Mục tiêu >= 90%)", f"{report.mobile_ratio_pct}%"])
    ws.append(["Tỷ lệ dòng có ít nhất 1 SĐT hợp lệ (%)", f"{report.valid_phone_ratio_pct}%"])
    ws.append([])

    # Phần 4: Đánh giá tiêu chí nghiệm thu
    add_section("4. ĐÁNH GIÁ TIÊU CHÍ NGHIỆM THU (BA §10.1)")
    ws.append(["Đủ sản lượng sàn tổng (>= 5.000 dòng)", "ĐẠT" if report.is_total_quota_passed else "CHƯA ĐẠT"])
    ws.append(["Đủ sản lượng sàn từng phân khúc", "ĐẠT" if report.is_segment_quota_passed else "CHƯA ĐẠT"])
    ws.append(["Tỷ lệ số di động đạt >= 90%", "ĐẠT" if report.is_mobile_ratio_passed else "CHƯA ĐẠT"])
    ws.append(["100% dòng có SĐT hợp lệ", "ĐẠT" if report.is_valid_phone_passed else "CHƯA ĐẠT"])
    ws.append(["KẾT LUẬN NGHIỆM THU CHUNG", "ĐẠT CHUẨN BÀN GIAO" if report.is_overall_passed else "CẦN BỔ SUNG"])

    res_cell = ws.cell(row=ws.max_row, column=2)
    res_cell.font = pass_font if report.is_overall_passed else fail_font

    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 15
